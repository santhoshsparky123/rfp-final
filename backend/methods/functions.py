from passlib.context import CryptContext
from fastapi import Depends, HTTPException
from fastapi import status
from typing import Optional, List
from datetime import datetime, timedelta
from fastapi.security.http import HTTPAuthorizationCredentials
from fastapi.security import HTTPBearer
from models.schema import User, UserRole, Company, SubscriptionStatus,Employee
from jose import JWTError, jwt  # Add this import for JWTError and jwt
import os
from dotenv import load_dotenv
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker, declarative_base, Session

import docx  # For Word
import openpyxl  # For Excel
from io import BytesIO

# Configuration
DATABASE_URL = os.getenv("DATABASE_URL")  
# if not DATABASE_URL:
#     raise ValueError("DATABASE_URL environment variable is not set")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 30

# Database setup
engine = create_engine(DATABASE_URL)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# Password hashing
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

# JWT Security
security = HTTPBearer()

load_dotenv()  # Load environment variables from .env file

SECRET_KEY = os.getenv("JWT_SECRET")

# Database dependency
def get_db():
    print("connected")
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# Utility functions
def verify_password(plain_password, hashed_password):
    return pwd_context.verify(plain_password, hashed_password)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict, expires_delta: Optional[timedelta] = None):
    to_encode = data.copy()
    if expires_delta:
        expire = datetime.utcnow() + expires_delta
    else:
        expire = datetime.utcnow() + timedelta(minutes=15)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

def authenticate_user(db: Session, username: str, password: str):
    user = db.query(User).filter(User.username == username).first()
    if not user:
        return False
    if not verify_password(password, user.hashed_password):
        return False
    return user

def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security), db: Session = Depends(get_db)):
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    user = db.query(User).filter(User.username == username).first()
    if user is None:
        raise credentials_exception
    return user  # <-- return the user object, not user.role

def require_role(required_roles: List[UserRole]):
    def role_checker(current_user: User = Depends(get_current_user)):
        if current_user is None or current_user.role not in required_roles:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Insufficient permissions"
            )
        return current_user
    return role_checker

def require_role1(required_roles: List[UserRole]):
    def role_checker1(current_user: Employee = Depends(get_current_user1)):
        if current_user.role not in required_roles:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Insufficient permissions"
            )
        return current_user
    return role_checker1

def get_current_user1(credentials: HTTPAuthorizationCredentials = Depends(security), db: Session = Depends(get_db)):
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    
    user = db.query(Employee).filter(Employee.name == username).first()
    if user is None:
        raise credentials_exception
    return user
    
    
    
def get_company_from_subdomain(subdomain: str, db: Session):
    company = db.query(Company).filter(Company.subdomain == subdomain).first()
    if not company:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Company not found"
        )
    if company.subscription_status != SubscriptionStatus.ACTIVE:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Company subscription expired"
        )
    return company


from langchain_community.document_loaders import PyPDFLoader
from typing import List

def extract_text_from_pdf(file_path: str) -> str:
    loader = PyPDFLoader(file_path)
    documents = loader.load()
    # Combine all page contents into a single string
    return "\n".join([doc.page_content for doc in documents])

import tempfile
def extract_text_from_pdf_bytes(file_bytes: bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes)
        tmp.flush()
        tmp_path = tmp.name

    try:
        text = extract_text_from_pdf(tmp_path)
    finally:
        os.remove(tmp_path)
    return text


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_excel(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    text_chunks = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = ' | '.join([str(cell) if cell is not None else '' for cell in row])
            text_chunks.append(line)
    return "\n".join(text_chunks)
